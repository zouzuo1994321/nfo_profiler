# -*- coding: utf-8 -*-
"""导出层：CSV / JSON / XLSX / Markdown。

CSV 用 UTF-8-BOM 编码，Excel 双击直接打开不乱码。
XLSX 需要 openpyxl（可选依赖，缺失时自动降级并提示）。
"""

from __future__ import annotations

import csv
import json
import os
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Sequence

from . import APP_NAME, __version__


# ---------------------------------------------------------------------------
# 作品明细
# ---------------------------------------------------------------------------

DETAIL_COLUMNS: Sequence[tuple] = (
    ("num", "番号"),
    ("title", "标题"),
    ("clean_title", "清洁标题"),
    ("year", "年份"),
    ("premiered", "发行日期"),
    ("dateadded", "入库时间"),
    ("studio", "片商"),
    ("maker", "制作商"),
    ("publisher", "发行商"),
    ("label", "厂牌"),
    ("series", "系列"),
    ("director", "导演"),
    ("actors", "演员"),
    ("tags", "内容标签"),
    ("tech_tags", "技术标签"),
    ("censor_status", "马赛克"),
    ("resolution", "分辨率"),
    ("video_codec", "视频编码"),
    ("runtime_min", "时长(分)"),
    ("userrating", "评分"),
    ("watched", "已看"),
    ("playcount", "播放次数"),
    ("num_prefix", "番号前缀"),
    ("original_filename", "视频文件"),
    ("video_size", "视频大小(字节)"),
    ("website", "来源链接"),
    ("path", "NFO路径"),
)


def iter_movie_rows(store, *, limit: int = 0, order_by: str = "userrating DESC, id",
                    source: Optional[str] = None):
    """流式产出作品明细（含聚合后的演员与标签字符串）。

    source 非空时只输出该数据源的记录。
    """
    sql = f"""
      SELECT m.*,
             (SELECT group_concat(t.tag, ' / ') FROM movie_tags t WHERE t.movie_id = m.id) AS tags,
             (SELECT group_concat(a.actor, ' / ') FROM movie_actors a WHERE a.movie_id = m.id) AS actors,
             (SELECT group_concat(d.tech, ' / ') FROM movie_tech d WHERE d.movie_id = m.id) AS tech_tags
      FROM movies m
    """
    params: List[Any] = []
    if source:
        sql += " WHERE m.source = ?"
        params.append(source)
    sql += f" ORDER BY {order_by}"
    if limit:
        sql += f" LIMIT {int(limit)}"
    for row in store.conn.execute(sql, params):
        d = dict(row)
        d["video_size_gb"] = round((d.get("video_size") or 0) / 1073741824, 2)
        yield d


def movie_rows(store, *, limit: int = 0, order_by: str = "userrating DESC, id",
               source: Optional[str] = None) -> List[Dict[str, Any]]:
    return list(iter_movie_rows(store, limit=limit, order_by=order_by, source=source))


def movies_for_report(store, limit: int = 1000, source: Optional[str] = None) -> List[Dict[str, Any]]:
    """报告内嵌的精简明细（只保留表格展示的列，控制 HTML 体积）。"""
    keys = [k for k, _ in DETAIL_COLUMNS if k not in ("path", "video_size", "website", "clean_title",
                                                      "maker", "publisher", "label", "tech_tags",
                                                      "video_codec", "playcount", "original_filename")]
    out = []
    for row in iter_movie_rows(store, limit=limit, source=source):
        item = {}
        for k in keys:
            v = row.get(k, "")
            if k == "title" and v and len(v) > 46:
                v = v[:46] + "…"
            if k == "actors" and v and len(v) > 40:
                v = v[:40] + "…"
            if k == "tags" and v and len(v) > 46:
                v = v[:46] + "…"
            item[k] = v
        out.append(item)
    return out


# ---------------------------------------------------------------------------
# 通用写文件
# ---------------------------------------------------------------------------

def _ensure_dir(path: str) -> None:
    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)


def _write_csv(path: str, header: Sequence[str], rows: Iterable[Sequence[Any]]) -> str:
    _ensure_dir(path)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        for r in rows:
            w.writerow(r)
    return path


def _freq_rows(items: Sequence[Dict[str, Any]], name_header: str = "名称") -> List[List[Any]]:
    out = []
    for i, it in enumerate(items or [], 1):
        aliases = "、".join(it.get("aliases") or [])
        out.append([i, it.get("name", ""), it.get("count", 0), aliases])
    return out


def _dist_rows(items: Sequence[Dict[str, Any]], category: str) -> List[List[Any]]:
    return [[category, it.get("name", ""), it.get("count", 0)] for it in (items or [])]


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def export_csv(
    data: Dict[str, Any],
    out_dir: str,
    *,
    store=None,
    include_movies: bool = True,
    movie_limit: int = 0,
    source: Optional[str] = None,
) -> List[str]:
    """导出一整套 CSV，返回生成的文件列表。"""
    os.makedirs(out_dir, exist_ok=True)
    written: List[str] = []

    ov = data.get("overview") or {}
    written.append(_write_csv(
        os.path.join(out_dir, "01_总览.csv"),
        ["指标", "值"],
        [
            ["收录作品数", ov.get("total", 0)],
            ["涉及艺人数", (ov.get("distinct") or {}).get("actors", 0)],
            ["内容标签数", (ov.get("distinct") or {}).get("tags", 0)],
            ["片商数", (ov.get("distinct") or {}).get("studios", 0)],
            ["系列数", (ov.get("distinct") or {}).get("series", 0)],
            ["导演数", (ov.get("distinct") or {}).get("directors", 0)],
            ["番号前缀数", (ov.get("distinct") or {}).get("prefixes", 0)],
            ["总时长(小时)", ov.get("total_hours", 0)],
            ["总时长(天)", ov.get("total_days", 0)],
            ["平均时长(分钟)", ov.get("avg_runtime", 0)],
            ["平均评分", ov.get("avg_rating", 0)],
            ["已评分作品数", ov.get("rated", 0)],
            ["已评分占比(%)", ov.get("rated_ratio", 0)],
            ["视频总容量(GB)", ov.get("total_gb", 0)],
            ["检测到视频文件数", ov.get("with_video", 0)],
            ["已观看数", ov.get("watched", 0)],
            ["播放总次数", ov.get("plays", 0)],
            ["平均标签数", ov.get("avg_tags", 0)],
            ["平均演员数", ov.get("avg_actors", 0)],
            ["年份范围", f"{ov.get('year_min','')}-{ov.get('year_max','')}"],
            ["生成时间", data.get("generated_at", "")],
        ],
    ))

    p = data.get("portrait") or {}
    if p and not p.get("empty"):
        written.append(_write_csv(
            os.path.join(out_dir, "02_用户画像.csv"),
            ["维度", "强度(0-100)", "说明"],
            [[r["name"], round(r["value"], 1), r.get("hint", "")] for r in p.get("radar", [])]
            + [["一句话画像", "", p.get("summary", "")]],
        ))

    freq_specs = [
        ("03_高频标签.csv", data.get("tags"), "标签"),
        ("04_技术标签.csv", data.get("tech_tags"), "技术标签"),
        ("05_高频艺人.csv", data.get("actors"), "艺人"),
        ("06_片商.csv", data.get("studios"), "片商"),
        ("07_发行商.csv", data.get("publishers"), "发行商"),
        ("08_厂牌.csv", data.get("labels"), "厂牌"),
        ("09_系列.csv", data.get("series"), "系列"),
        ("10_导演.csv", data.get("directors"), "导演"),
        ("11_番号前缀.csv", data.get("prefixes"), "番号前缀"),
        ("12_剧情高频词.csv", data.get("keywords"), "关键词"),
    ]
    for fname, items, label in freq_specs:
        if not items:
            continue
        written.append(_write_csv(os.path.join(out_dir, fname),
                                  ["排名", label, "出现次数", "其他写法"],
                                  _freq_rows(items, label)))

    # 艺人明细（含代表片商/标签）
    if data.get("actors"):
        rows = []
        for i, a in enumerate(data["actors"], 1):
            rows.append([
                i, a.get("name", ""), a.get("count", 0), a.get("avg_rating", ""),
                f"{a.get('year_from','')}-{a.get('year_to','')}",
                a.get("avg_runtime", ""), a.get("watched", 0),
                "、".join(s["name"] for s in a.get("top_studios", [])),
                "、".join(t["name"] for t in a.get("top_tags", [])),
                "、".join(x["name"] for x in a.get("top_partners", [])),
            ])
        written.append(_write_csv(
            os.path.join(out_dir, "13_艺人明细.csv"),
            ["排名", "艺人", "作品数", "平均评分", "活跃年份", "平均时长", "已看",
             "代表片商", "代表标签", "常搭档"], rows))

    # 分布
    dist_rows: List[List[Any]] = []
    for key, label in [
        ("dist_resolution", "分辨率"), ("dist_censor", "马赛克"), ("dist_rating", "评分"),
        ("dist_runtime", "时长"), ("dist_year", "发行年份"), ("dist_added_month", "入库月份"),
        ("dist_added_hour", "入库时段"), ("dist_weekday", "入库星期"),
        ("dist_actor_count", "演出人数"),
    ]:
        dist_rows.extend(_dist_rows(data.get(key), label))
    if dist_rows:
        written.append(_write_csv(os.path.join(out_dir, "14_分布统计.csv"),
                                  ["维度", "分类", "数量"], dist_rows))

    # 共现 / 共演
    co = data.get("cooccurrence") or {}
    if co.get("edges"):
        written.append(_write_csv(
            os.path.join(out_dir, "15_标签共现.csv"), ["排名", "标签A", "标签B", "共现次数"],
            [[i, e["source"], e["target"], e["value"]] for i, e in enumerate(co["edges"], 1)]))
    if data.get("actor_pairs"):
        written.append(_write_csv(
            os.path.join(out_dir, "16_共演组合.csv"), ["排名", "艺人A", "艺人B", "合作次数"],
            [[i, p["a"], p["b"], p["count"]] for i, p in enumerate(data["actor_pairs"], 1)]))

    # 作品明细
    if include_movies and store is not None:
        path = os.path.join(out_dir, "17_作品明细.csv")
        _ensure_dir(path)
        header = [label for _, label in DETAIL_COLUMNS] + ["视频大小(GB)"]
        keys = [k for k, _ in DETAIL_COLUMNS]
        with open(path, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(header)
            n = 0
            for row in iter_movie_rows(store, limit=movie_limit, source=source):
                w.writerow([row.get(k, "") for k in keys] + [row.get("video_size_gb", "")])
                n += 1
        written.append(path)

    return written


# ---------------------------------------------------------------------------
# JSON
# ---------------------------------------------------------------------------

def export_json(data: Dict[str, Any], path: str) -> str:
    _ensure_dir(path)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
    return path


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def export_xlsx(
    data: Dict[str, Any],
    path: str,
    *,
    store=None,
    include_movies: bool = True,
    movie_limit: int = 50000,
    source: Optional[str] = None,
) -> Optional[str]:
    """导出多工作表 XLSX。缺少 openpyxl 时返回 None。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    _ensure_dir(path)
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="2F5597")
    head_font = Font(color="FFFFFF", bold=True, size=11)

    def add_sheet(name: str, header: Sequence[str], rows: Iterable[Sequence[Any]],
                  widths: Optional[Sequence[int]] = None):
        ws = wb.create_sheet(name[:31])
        ws.append(list(header))
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        n = 0
        for r in rows:
            ws.append(list(r))
            n += 1
        ws.freeze_panes = "A2"
        if widths:
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
        else:
            for i in range(1, len(header) + 1):
                ws.column_dimensions[get_column_letter(i)].width = 16
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(n + 1, 2)}"
        return ws

    wb.remove(wb.active)

    ov = data.get("overview") or {}
    dist = ov.get("distinct") or {}
    add_sheet("总览", ["指标", "值"], [
        ["收录作品数", ov.get("total", 0)],
        ["涉及艺人数", dist.get("actors", 0)],
        ["内容标签数", dist.get("tags", 0)],
        ["片商数", dist.get("studios", 0)],
        ["系列数", dist.get("series", 0)],
        ["导演数", dist.get("directors", 0)],
        ["总时长(小时)", ov.get("total_hours", 0)],
        ["平均时长(分钟)", ov.get("avg_runtime", 0)],
        ["平均评分", ov.get("avg_rating", 0)],
        ["已评分占比(%)", ov.get("rated_ratio", 0)],
        ["视频总容量(GB)", ov.get("total_gb", 0)],
        ["年份范围", f"{ov.get('year_min','')}-{ov.get('year_max','')}"],
        ["生成时间", data.get("generated_at", "")],
    ], widths=[22, 20])

    p = data.get("portrait") or {}
    if p and not p.get("empty"):
        add_sheet("用户画像", ["维度", "强度(0-100)", "说明"],
                  [[r["name"], round(r["value"], 1), r.get("hint", "")] for r in p.get("radar", [])]
                  + [[("一句话画像"), "", p.get("summary", "")]], widths=[16, 14, 46])

    for sheet, items, label in [
        ("高频标签", data.get("tags"), "标签"),
        ("技术标签", data.get("tech_tags"), "技术标签"),
        ("高频艺人", data.get("actors"), "艺人"),
        ("片商", data.get("studios"), "片商"),
        ("发行商", data.get("publishers"), "发行商"),
        ("厂牌", data.get("labels"), "厂牌"),
        ("系列", data.get("series"), "系列"),
        ("导演", data.get("directors"), "导演"),
        ("番号前缀", data.get("prefixes"), "番号前缀"),
        ("剧情高频词", data.get("keywords"), "关键词"),
    ]:
        if items:
            add_sheet(sheet, ["排名", label, "出现次数", "其他写法"], _freq_rows(items, label),
                      widths=[8, 30, 14, 26])

    if data.get("actors"):
        add_sheet("艺人明细",
                  ["排名", "艺人", "作品数", "平均评分", "起始年", "结束年", "平均时长",
                   "已看", "代表片商", "代表标签", "常搭档"],
                  [[i, a.get("name", ""), a.get("count", 0), a.get("avg_rating", ""),
                    a.get("year_from", ""), a.get("year_to", ""), a.get("avg_runtime", ""),
                    a.get("watched", 0),
                    "、".join(s["name"] for s in a.get("top_studios", [])),
                    "、".join(t["name"] for t in a.get("top_tags", [])),
                    "、".join(x["name"] for x in a.get("top_partners", []))]
                   for i, a in enumerate(data["actors"], 1)],
                  widths=[8, 20, 10, 10, 10, 10, 11, 8, 26, 34, 24])

    dist_rows: List[List[Any]] = []
    for key, label in [
        ("dist_resolution", "分辨率"), ("dist_censor", "马赛克"), ("dist_rating", "评分"),
        ("dist_runtime", "时长"), ("dist_year", "发行年份"), ("dist_added_month", "入库月份"),
        ("dist_added_hour", "入库时段"), ("dist_weekday", "入库星期"),
        ("dist_actor_count", "演出人数"),
    ]:
        dist_rows.extend(_dist_rows(data.get(key), label))
    if dist_rows:
        add_sheet("分布统计", ["维度", "分类", "数量"], dist_rows, widths=[16, 18, 12])

    co = data.get("cooccurrence") or {}
    if co.get("edges"):
        add_sheet("标签共现", ["排名", "标签A", "标签B", "共现次数"],
                  [[i, e["source"], e["target"], e["value"]] for i, e in enumerate(co["edges"], 1)],
                  widths=[8, 22, 22, 12])
    if data.get("actor_pairs"):
        add_sheet("共演组合", ["排名", "艺人A", "艺人B", "合作次数"],
                  [[i, x["a"], x["b"], x["count"]] for i, x in enumerate(data["actor_pairs"], 1)],
                  widths=[8, 22, 22, 12])

    if include_movies and store is not None:
        header = [label for _, label in DETAIL_COLUMNS] + ["视频大小(GB)"]
        keys = [k for k, _ in DETAIL_COLUMNS]
        ws = wb.create_sheet("作品明细")
        ws.append(header)
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
        n = 0
        for row in iter_movie_rows(store, limit=movie_limit, source=source):
            ws.append([row.get(k, "") for k in keys] + [row.get("video_size_gb", "")])
            n += 1
        ws.freeze_panes = "A2"
        for i, w in enumerate([14, 52, 52, 8, 12, 17, 18, 14, 14, 14, 34, 14, 30, 60, 24, 10, 11,
                               12, 10, 8, 8, 10, 12, 22, 16, 34, 46, 12], 1):
            if i <= len(header):
                ws.column_dimensions[get_column_letter(i)].width = w
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(n + 1, 2)}"

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def export_markdown(data: Dict[str, Any], out_dir_or_path: str) -> str:
    """生成一份可读的 Markdown 画像卡片。"""
    ov = data.get("overview") or {}
    p = data.get("portrait") or {}
    if os.path.isdir(out_dir_or_path) or out_dir_or_path.endswith(("/", "\\")):
        path = os.path.join(out_dir_or_path, "用户画像.md")
    else:
        path = out_dir_or_path
    _ensure_dir(path)

    L: List[str] = []
    A = L.append
    A(f"# 用户×癖分析报告\n")
    A(f"> 生成时间：{data.get('generated_at','')}　|　工具：{APP_NAME} v{__version__}\n")

    if not ov.get("total"):
        A("\n数据库中没有记录。\n")
    else:
        dist = ov.get("distinct") or {}
        A("\n## 总览\n")
        A("| 指标 | 值 | 指标 | 值 |")
        A("|---|---|---|---|")
        A(f"| 收录作品 | {ov.get('total',0):,} 部 | 涉及艺人 | {dist.get('actors',0):,} 位 |")
        A(f"| 内容标签 | {dist.get('tags',0):,} 个 | 片商 | {dist.get('studios',0):,} 家 |")
        A(f"| 系列 | {dist.get('series',0):,} 个 | 导演 | {dist.get('directors',0):,} 位 |")
        A(f"| 总时长 | {ov.get('total_hours',0):,} 小时 | 平均时长 | {ov.get('avg_runtime',0)} 分钟 |")
        A(f"| 平均评分 | {ov.get('avg_rating',0)} 分 | 已评分占比 | {ov.get('rated_ratio',0)}% |")
        A(f"| 视频容量 | {ov.get('total_gb',0):,} GB | 年份跨度 | {ov.get('year_min','')}–{ov.get('year_max','')} |")

        if p and not p.get("empty"):
            A("\n## 一句话画像\n")
            A(f"> {p.get('summary','')}\n")
            A("\n## 偏好雷达\n")
            A("| 维度 | 强度 | 说明 |")
            A("|---|---|---|")
            for r in p.get("radar", []):
                A(f"| {r['name']} | {r['value']:.0f} / 100 | {r.get('hint','')} |")

        def table(title: str, items, label: str, extra=None):
            if not items:
                return
            A(f"\n## {title}\n")
            if extra:
                A("| 排名 | " + label + " | " + " | ".join(extra[0]) + " |")
            else:
                A("| 排名 | " + label + " | 出现次数 |")
            A("|---" * (len(extra[0]) + 2 if extra else 3) + "|")
            for i, it in enumerate(items, 1):
                if extra:
                    vals = " | ".join(str(f(it)) for f in extra[1])
                    A(f"| {i} | {it.get('name','')} | {vals} |")
                else:
                    A(f"| {i} | {it.get('name','')} | {it.get('count',0):,} |")

        table("高频标签 Top 20", (data.get("tags") or [])[:20], "标签")
        table("高频艺人 Top 20", (data.get("actors") or [])[:20], "艺人",
              (["作品数", "平均评分", "活跃年份"],
               [lambda a: a.get("count", 0), lambda a: a.get("avg_rating", ""),
                lambda a: f"{a.get('year_from','')}–{a.get('year_to','')}"]))
        table("片商 Top 10", (data.get("studios") or [])[:10], "片商")
        table("系列 Top 10", (data.get("series") or [])[:10], "系列")
        table("导演 Top 10", (data.get("directors") or [])[:10], "导演")
        table("剧情高频词 Top 20", (data.get("keywords") or [])[:20], "关键词")

        for key, label in [("dist_resolution", "画质分布"), ("dist_censor", "马赛克分布"),
                           ("dist_runtime", "时长分布")]:
            items = [d for d in (data.get(key) or []) if d.get("count")]
            if items:
                total = sum(d["count"] for d in items) or 1
                A(f"\n## {label}\n")
                A("| 分类 | 数量 | 占比 |")
                A("|---|---|---|")
                for d in items:
                    A(f"| {d['name']} | {d['count']:,} | {d['count']*100/total:.1f}% |")

    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))
    return path


def export_all(
    data: Dict[str, Any],
    out_dir: str,
    *,
    store=None,
    formats: Sequence[str] = ("csv", "json", "xlsx", "md"),
    html: bool = True,
    movie_limit: int = 0,
    report_movies: int = 1000,
    source_label: str = "",
    source: Optional[str] = None,
) -> Dict[str, Any]:
    """一次性导出所有格式，返回 {格式: 路径或列表}。"""
    from .report import write_html_report

    os.makedirs(out_dir, exist_ok=True)
    result: Dict[str, Any] = {}

    if "json" in formats:
        result["json"] = export_json(data, os.path.join(out_dir, "画像数据.json"))
    if "csv" in formats:
        result["csv"] = export_csv(data, os.path.join(out_dir, "csv"),
                                   store=store, movie_limit=movie_limit, source=source)
    if "xlsx" in formats:
        p = export_xlsx(data, os.path.join(out_dir, "画像数据.xlsx"), store=store, source=source)
        result["xlsx"] = p or "未生成（缺少 openpyxl，可 pip install openpyxl）"
    if "md" in formats:
        result["md"] = export_markdown(data, out_dir)
    if html:
        movies = movies_for_report(store, limit=report_movies, source=source) if store else []
        result["html"] = write_html_report(
            data, os.path.join(out_dir, "用户画像报告.html"),
            movies=movies, source_label=source_label,
        )
    return result


__all__ = [
    "export_csv", "export_json", "export_xlsx", "export_markdown", "export_all",
    "iter_movie_rows", "movie_rows", "movies_for_report", "DETAIL_COLUMNS",
]
